from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from threading import RLock, Thread
from time import sleep
from typing import Any, Iterable

from app.core.config import get_settings
from app.services.customer_service import normalize_phone
from app.services.demo_store import DEMO_DISPATCH_IMPORT_BATCHES, DEMO_DISPATCH_IMPORT_ITEMS, new_id, persist_demo_state
from app.services.dispatch_service import create_dispatch, get_dispatch, retry_dispatch

MAX_UPLOAD_BYTES = 12 * 1024 * 1024
MAX_IMPORTED_ROWS = 20000
SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xls", ".ods", ".pdf", ".docx"}
NAME_ALIASES = {
    "nome", "nome cliente", "nome do cliente", "cliente", "contato", "nome contato",
    "nome do contato", "name", "customer", "customer name", "contact", "contact name",
}
PHONE_ALIASES = {
    "telefone", "telefone cliente", "telefone do cliente", "fone", "celular", "whatsapp",
    "numero", "numero telefone", "numero do telefone", "phone", "mobile", "cell", "telephone",
}
IMPORT_STATUSES = {"ready", "invalid", "duplicate", "pending", "submitted", "sent", "delivered", "read", "failed", "cancelled"}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _norm_header(value: Any) -> str:
    text = str(value or "").strip().casefold()
    text = "".join(ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch))
    text = re.sub(r"[^a-z0-9]+", " ", text).strip()
    return re.sub(r"\s+", " ", text)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def _valid_phone(value: Any) -> str | None:
    raw = _cell_text(value)
    digits = normalize_phone(raw)
    if not digits:
        return None
    # Meta/WhatsApp normally needs country + area code + subscriber. Keep this
    # permissive enough for international lists while blocking obvious IDs.
    if len(digits) < 10 or len(digits) > 15:
        return None
    return digits


def _decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _rows_csv(data: bytes) -> list[list[Any]]:
    text = _decode_csv(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        if ";" in sample:
            dialect.delimiter = ";"
    return [list(row) for row in csv.reader(io.StringIO(text), dialect)]


def _rows_xlsx(data: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows: list[list[Any]] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) >= MAX_IMPORTED_ROWS + 30:
                return rows
    return rows


def _rows_xls(data: bytes) -> list[list[Any]]:
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError("Para importar XLS, reinicie o VENDRA para instalar as dependências da versão atual.") from exc
    book = xlrd.open_workbook(file_contents=data)
    rows: list[list[Any]] = []
    for sheet in book.sheets():
        for r in range(sheet.nrows):
            rows.append([sheet.cell_value(r, c) for c in range(sheet.ncols)])
            if len(rows) >= MAX_IMPORTED_ROWS + 30:
                return rows
    return rows


def _rows_ods(data: bytes) -> list[list[Any]]:
    from odf.opendocument import load
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = load(io.BytesIO(data))
    rows: list[list[Any]] = []
    for table in doc.spreadsheet.getElementsByType(Table):
        for tr in table.getElementsByType(TableRow):
            current: list[str] = []
            for cell in tr.getElementsByType(TableCell):
                text_parts = []
                for paragraph in cell.getElementsByType(P):
                    for child in paragraph.childNodes:
                        if hasattr(child, "data"):
                            text_parts.append(str(child.data))
                value = " ".join(text_parts).strip()
                repeated = int(cell.getAttribute("numbercolumnsrepeated") or 1)
                current.extend([value] * min(repeated, 50))
            rows.append(current)
            if len(rows) >= MAX_IMPORTED_ROWS + 30:
                return rows
    return rows


def _rows_docx(data: bytes) -> list[list[Any]]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    rows: list[list[Any]] = []
    for table in doc.tables:
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
            if len(rows) >= MAX_IMPORTED_ROWS + 30:
                return rows
    return rows


def _rows_pdf(data: bytes) -> tuple[list[list[Any]], list[str]]:
    import pdfplumber

    rows: list[list[Any]] = []
    lines: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            try:
                for table in page.extract_tables() or []:
                    for row in table or []:
                        rows.append(list(row or []))
                        if len(rows) >= MAX_IMPORTED_ROWS + 30:
                            return rows, lines
            except Exception:
                pass
            text = page.extract_text() or ""
            lines.extend(line.strip() for line in text.splitlines() if line.strip())
    return rows, lines


def _header_indices(rows: list[list[Any]]) -> tuple[int | None, int | None, int | None]:
    for row_index, row in enumerate(rows[:20]):
        name_index = None
        phone_index = None
        for idx, value in enumerate(row):
            key = _norm_header(value)
            if key in NAME_ALIASES and name_index is None:
                name_index = idx
            if key in PHONE_ALIASES and phone_index is None:
                phone_index = idx
        if name_index is not None and phone_index is not None:
            return row_index, name_index, phone_index
    return None, None, None


def _infer_indices(rows: list[list[Any]]) -> tuple[int | None, int | None]:
    sample = [row for row in rows[:80] if row]
    max_cols = max((len(row) for row in sample), default=0)
    if max_cols == 0:
        return None, None
    phone_scores: list[tuple[int, int]] = []
    text_scores: list[tuple[int, int]] = []
    for col in range(max_cols):
        phones = 0
        texts = 0
        for row in sample:
            value = row[col] if col < len(row) else None
            raw = _cell_text(value)
            if _valid_phone(value):
                phones += 1
            elif raw and any(ch.isalpha() for ch in raw) and len(raw) <= 180:
                texts += 1
        phone_scores.append((phones, col))
        text_scores.append((texts, col))
    phone_score, phone_col = max(phone_scores, default=(0, -1))
    if phone_score < max(2, len(sample) // 5):
        return None, None
    possible_names = [(score, col) for score, col in text_scores if col != phone_col]
    name_score, name_col = max(possible_names, default=(0, -1))
    if name_score == 0:
        return None, phone_col
    return name_col, phone_col


def _extract_records_from_rows(rows: list[list[Any]]) -> list[dict[str, Any]]:
    rows = [list(row) for row in rows if any(_cell_text(value) for value in row)]
    if not rows:
        return []
    header_row, name_col, phone_col = _header_indices(rows)
    start = (header_row + 1) if header_row is not None else 0
    if name_col is None or phone_col is None:
        name_col, phone_col = _infer_indices(rows[start:])
    if phone_col is None:
        return []

    records: list[dict[str, Any]] = []
    for idx, row in enumerate(rows[start:], start=start + 1):
        if len(records) >= MAX_IMPORTED_ROWS:
            break
        phone_raw = _cell_text(row[phone_col] if phone_col < len(row) else "")
        name = _cell_text(row[name_col] if name_col is not None and name_col < len(row) else "")
        if not name and phone_raw and name_col is None:
            # In headerless two-column exports, use the first plausible text cell.
            name = next((_cell_text(v) for c, v in enumerate(row) if c != phone_col and any(ch.isalpha() for ch in _cell_text(v))), "")
        if not name and not phone_raw:
            continue
        records.append({"row_number": idx, "name": name[:160], "phone_raw": phone_raw[:80]})
    return records


_PHONE_IN_LINE = re.compile(r"(?<!\d)(?:\+?\d[\d\s().-]{8,}\d)(?!\d)")


def _extract_records_from_pdf_lines(lines: Iterable[str]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for idx, line in enumerate(lines, start=1):
        if len(records) >= MAX_IMPORTED_ROWS:
            break
        match = _PHONE_IN_LINE.search(line)
        if not match:
            continue
        phone_raw = match.group(0).strip()
        name = (line[: match.start()] + " " + line[match.end() :]).strip(" -–—|;,:\t")
        name = re.sub(r"\s+", " ", name)
        if not name or _norm_header(name) in PHONE_ALIASES:
            continue
        records.append({"row_number": idx, "name": name[:160], "phone_raw": phone_raw[:80]})
    return records


def parse_contact_file(*, filename: str, content: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    filename = Path(filename or "lista").name
    ext = Path(filename).suffix.casefold()
    if ext not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Formato não suportado. Use arquivos {supported}.")
    if not content:
        raise ValueError("O arquivo está vazio.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("O arquivo ultrapassa o limite de 12 MB por importação.")

    warnings: list[str] = []
    pdf_lines: list[str] = []
    try:
        if ext == ".csv":
            rows = _rows_csv(content)
        elif ext == ".xlsx":
            rows = _rows_xlsx(content)
        elif ext == ".xls":
            rows = _rows_xls(content)
        elif ext == ".ods":
            rows = _rows_ods(content)
        elif ext == ".docx":
            rows = _rows_docx(content)
        else:
            rows, pdf_lines = _rows_pdf(content)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Não foi possível ler o arquivo {ext or filename}: {str(exc)[:220]}") from exc

    records = _extract_records_from_rows(rows)
    if not records and ext == ".pdf":
        records = _extract_records_from_pdf_lines(pdf_lines)
        if records:
            warnings.append("O PDF não expôs uma tabela estruturada; os contatos foram identificados pelo texto da página.")
    if not records:
        raise ValueError("Não foi possível identificar colunas/linhas com nome e telefone nesse arquivo.")
    if len(records) >= MAX_IMPORTED_ROWS:
        warnings.append(f"A importação foi limitada aos primeiros {MAX_IMPORTED_ROWS} contatos identificados.")
    return records, warnings


def import_contact_file(*, filename: str, content: bytes, imported_by: str | None) -> dict[str, Any]:
    if not get_settings().demo_mode:
        from app.services.dispatch_list_service import import_file
        return import_file(filename=filename,content=content,imported_by=imported_by)
    records, warnings = parse_contact_file(filename=filename, content=content)
    batch_id = new_id("dispatch-list")
    now = _now()
    seen: set[str] = set()
    valid = invalid = duplicates = 0
    items: list[dict[str, Any]] = []
    for record in records:
        name = str(record.get("name") or "").strip()[:160]
        phone_raw = str(record.get("phone_raw") or "").strip()[:80]
        phone = _valid_phone(phone_raw)
        if not name or not phone:
            status = "invalid"
            error = "Nome ausente." if not name else "Telefone inválido; informe DDI, DDD e número."
            invalid += 1
        elif phone in seen:
            status = "duplicate"
            error = "Telefone repetido nesta lista."
            duplicates += 1
        else:
            status = "ready"
            error = None
            valid += 1
            seen.add(phone)
        items.append(
            {
                "id": new_id("dispatch-contact"),
                "batch_id": batch_id,
                "row_number": record.get("row_number"),
                "name": name,
                "phone_raw": phone_raw,
                "phone": phone,
                "status": status,
                "last_error": error,
                "dispatch_id": None,
                "created_at": now,
                "updated_at": now,
            }
        )
    batch = {
        "id": batch_id,
        "filename": Path(filename or "lista").name[:220],
        "extension": Path(filename or "").suffix.casefold(),
        "total_count": len(items),
        "valid_count": valid,
        "invalid_count": invalid,
        "duplicate_count": duplicates,
        "warnings": warnings,
        "imported_by": imported_by,
        "created_at": now,
        "updated_at": now,
    }
    DEMO_DISPATCH_IMPORT_BATCHES.append(batch)
    DEMO_DISPATCH_IMPORT_ITEMS.extend(items)
    return dict(batch)


def list_import_batches(limit: int = 30) -> list[dict[str, Any]]:
    if not get_settings().demo_mode:
        from app.services.dispatch_list_service import batches
        return batches(limit)
    rows = sorted(DEMO_DISPATCH_IMPORT_BATCHES, key=lambda row: row.get("created_at") or "", reverse=True)
    return [dict(row) for row in rows[: max(1, min(limit, 100))]]


def get_import_batch(batch_id: str | None) -> dict[str, Any] | None:
    if not get_settings().demo_mode:
        from app.services.dispatch_list_service import batch
        return batch(batch_id)
    if not batch_id:
        return None
    row = next((row for row in DEMO_DISPATCH_IMPORT_BATCHES if str(row.get("id")) == str(batch_id)), None)
    return dict(row) if row else None


def _status_label(status: str) -> str:
    labels = {
        "ready": "Pronto",
        "queued": "Na fila",
        "processing": "Reservado",
        "invalid": "Inválido",
        "duplicate": "Duplicado",
        "pending": "Pendente",
        "submitted": "Enviado ao Make",
        "sent": "Enviado",
        "delivered": "Entregue",
        "read": "Lido",
        "failed": "Falhou",
        "cancelled": "Cancelado",
    }
    return labels.get(status, status)


def list_import_items(batch_id: str) -> list[dict[str, Any]]:
    if not get_settings().demo_mode:
        from app.services.dispatch_list_service import items
        return items(batch_id)
    rows = [row for row in DEMO_DISPATCH_IMPORT_ITEMS if str(row.get("batch_id")) == str(batch_id)]
    rows.sort(key=lambda row: (int(row.get("row_number") or 0), str(row.get("id") or "")))
    public: list[dict[str, Any]] = []
    for source in rows:
        row = dict(source)
        dispatch_id = str(row.get("dispatch_id") or "")
        if dispatch_id:
            dispatch = get_dispatch(dispatch_id)
            if dispatch:
                row["status"] = dispatch.get("status") or row.get("status")
                row["last_error"] = dispatch.get("last_error") or row.get("last_error")
                row["provider_message_id"] = dispatch.get("provider_message_id")
        row["status_label"] = _status_label(str(row.get("status") or "ready"))
        public.append(row)
    return public



_BULK_LOCK = RLock()
_BULK_WORKERS: dict[str, Thread] = {}


def _bulk_public(batch: dict[str, Any] | None) -> dict[str, Any]:
    if not batch:
        return {
            "status": "idle",
            "total": 0,
            "processed": 0,
            "failed": 0,
            "remaining": 0,
            "percent": 0,
        }
    total = int(batch.get("bulk_total") or 0)
    processed = int(batch.get("bulk_processed") or 0)
    failed = int(batch.get("bulk_failed") or 0)
    status = str(batch.get("bulk_status") or "idle")
    percent = int(round((processed / total) * 100)) if total else (100 if status in {"completed", "completed_with_errors"} else 0)
    return {
        "status": status,
        "total": total,
        "processed": processed,
        "failed": failed,
        "remaining": max(total - processed, 0),
        "percent": max(0, min(percent, 100)),
        "started_at": batch.get("bulk_started_at"),
        "completed_at": batch.get("bulk_completed_at"),
        "current_item_id": batch.get("bulk_current_item_id"),
        "last_error": batch.get("bulk_last_error"),
    }


def get_bulk_dispatch_state(batch_id: str) -> dict[str, Any]:
    if not get_settings().demo_mode:
        from app.services.dispatch_list_service import state
        return state(batch_id)
    batch = next((row for row in DEMO_DISPATCH_IMPORT_BATCHES if str(row.get("id")) == str(batch_id)), None)
    if not batch:
        raise ValueError("Lista importada não encontrada.")
    # Se uma execução ficou marcada como ativa mas o processo anterior já morreu,
    # deixa claro que é possível retomá-la sem duplicar itens já enviados.
    status = str(batch.get("bulk_status") or "idle")
    worker = _BULK_WORKERS.get(str(batch_id))
    if status in {"queued", "running"} and (not worker or not worker.is_alive()):
        batch["bulk_status"] = "interrupted"
        batch["bulk_last_error"] = "A fila foi interrompida antes da conclusão. Clique em Enviar todos para continuar os contatos ainda prontos."
        batch["updated_at"] = _now()
        persist_demo_state()
    return _bulk_public(batch)


def _bulk_worker(batch_id: str, item_ids: list[str], requested_by: str | None) -> None:
    batch = next((row for row in DEMO_DISPATCH_IMPORT_BATCHES if str(row.get("id")) == str(batch_id)), None)
    if not batch:
        return
    try:
        batch["bulk_status"] = "running"
        batch["bulk_started_at"] = batch.get("bulk_started_at") or _now()
        batch["bulk_completed_at"] = None
        batch["bulk_last_error"] = None
        persist_demo_state()
        for item_id in item_ids:
            item = next((row for row in DEMO_DISPATCH_IMPORT_ITEMS if str(row.get("id")) == str(item_id)), None)
            if not item or str(item.get("status") or "") != "ready":
                # Pode ter sido disparado manualmente depois do clique no lote.
                batch["bulk_processed"] = int(batch.get("bulk_processed") or 0) + 1
                batch["updated_at"] = _now()
                persist_demo_state()
                continue
            batch["bulk_current_item_id"] = item_id
            batch["updated_at"] = _now()
            persist_demo_state()
            try:
                result = dispatch_import_item(item_id=item_id, requested_by=requested_by)
                dispatch = result.get("dispatch") or {}
                from app.services.audit_service import record_audit
                record_audit(
                    actor_id=requested_by,
                    entity_type="dispatch_import_item",
                    entity_id=item_id,
                    action="dispatch_requested_from_list_bulk",
                    payload={
                        "batch_id": batch_id,
                        "dispatch_id": dispatch.get("id"),
                        "phone": dispatch.get("phone"),
                        "status": dispatch.get("status"),
                    },
                )
                if str(dispatch.get("status") or "") == "failed":
                    batch["bulk_failed"] = int(batch.get("bulk_failed") or 0) + 1
            except Exception as exc:
                batch["bulk_failed"] = int(batch.get("bulk_failed") or 0) + 1
                batch["bulk_last_error"] = str(exc)[:1000]
            finally:
                batch["bulk_processed"] = int(batch.get("bulk_processed") or 0) + 1
                batch["updated_at"] = _now()
                persist_demo_state()
            # O processamento é propositalmente serial. Este pequeno intervalo
            # também evita uma rajada local caso o webhook do Make responda instantaneamente.
            sleep(0.15)
        batch["bulk_current_item_id"] = None
        batch["bulk_completed_at"] = _now()
        batch["bulk_status"] = "completed_with_errors" if int(batch.get("bulk_failed") or 0) else "completed"
        batch["updated_at"] = _now()
        from app.services.audit_service import record_audit
        record_audit(
            actor_id=requested_by,
            entity_type="dispatch_import_batch",
            entity_id=batch_id,
            action="bulk_dispatch_completed",
            payload={
                "total": batch.get("bulk_total"),
                "processed": batch.get("bulk_processed"),
                "failed": batch.get("bulk_failed"),
                "status": batch.get("bulk_status"),
            },
        )
        persist_demo_state()
    except Exception as exc:
        batch["bulk_status"] = "interrupted"
        batch["bulk_current_item_id"] = None
        batch["bulk_last_error"] = str(exc)[:1000]
        batch["updated_at"] = _now()
        persist_demo_state()
    finally:
        with _BULK_LOCK:
            _BULK_WORKERS.pop(str(batch_id), None)


def start_bulk_dispatch(*, batch_id: str, requested_by: str | None, quantity: int | None = None) -> dict[str, Any]:
    if not get_settings().demo_mode:
        from app.services.dispatch_list_service import start
        return start(batch_id=batch_id,requested_by=requested_by,quantity=quantity)
    batch = next((row for row in DEMO_DISPATCH_IMPORT_BATCHES if str(row.get("id")) == str(batch_id)), None)
    if not batch:
        raise ValueError("Lista importada não encontrada.")
    ready_ids = [
        str(row.get("id"))
        for row in DEMO_DISPATCH_IMPORT_ITEMS
        if str(row.get("batch_id")) == str(batch_id) and str(row.get("status") or "") == "ready"
    ]
    if not ready_ids:
        raise ValueError("Não existem contatos prontos para envio nesta lista.")
    if quantity is not None:
        if quantity < 1 or quantity > len(ready_ids):
            raise ValueError('Informe uma quantidade entre 1 e o total de contatos prontos.')
        ready_ids = ready_ids[:quantity]
    with _BULK_LOCK:
        active = _BULK_WORKERS.get(str(batch_id))
        if active and active.is_alive():
            return _bulk_public(batch)
        batch["bulk_status"] = "queued"
        batch["bulk_total"] = len(ready_ids)
        batch["bulk_processed"] = 0
        batch["bulk_failed"] = 0
        batch["bulk_current_item_id"] = None
        batch["bulk_started_at"] = _now()
        batch["bulk_completed_at"] = None
        batch["bulk_requested_by"] = requested_by
        batch["bulk_last_error"] = None
        batch["updated_at"] = _now()
        persist_demo_state()
        worker = Thread(
            target=_bulk_worker,
            args=(str(batch_id), ready_ids, requested_by),
            name=f"vendra-dispatch-bulk-{batch_id}",
            daemon=True,
        )
        _BULK_WORKERS[str(batch_id)] = worker
        worker.start()
    return _bulk_public(batch)

def dispatch_import_item(*, item_id: str, requested_by: str | None, batch_id: str | None = None) -> dict[str, Any]:
    if not get_settings().demo_mode:
        from app.services.dispatch_list_service import send_item
        return send_item(item_id=item_id,requested_by=requested_by,batch_id=batch_id)
    item = next((row for row in DEMO_DISPATCH_IMPORT_ITEMS if str(row.get("id")) == str(item_id)), None)
    if not item:
        raise ValueError("Contato importado não encontrado.")
    if item.get("status") in {"invalid", "duplicate"}:
        raise ValueError(item.get("last_error") or "Este contato não está apto para disparo.")

    dispatch_id = str(item.get("dispatch_id") or "")
    if dispatch_id:
        existing = get_dispatch(dispatch_id)
        if existing and existing.get("status") in {"failed", "cancelled"}:
            dispatch = retry_dispatch(dispatch_id, requested_by=requested_by)
        elif existing:
            return {"item": dict(item), "dispatch": existing, "already_requested": True}
        else:
            item["dispatch_id"] = None
            dispatch = create_dispatch(name=str(item.get("name") or ""), phone=str(item.get("phone") or item.get("phone_raw") or ""), requested_by=requested_by)
    else:
        dispatch = create_dispatch(name=str(item.get("name") or ""), phone=str(item.get("phone") or item.get("phone_raw") or ""), requested_by=requested_by)
        item["dispatch_id"] = str(dispatch.get("id") or "") or None

    item["status"] = dispatch.get("status") or "pending"
    item["last_error"] = dispatch.get("last_error")
    item["updated_at"] = _now()
    return {"item": dict(item), "dispatch": dispatch, "already_requested": False}
